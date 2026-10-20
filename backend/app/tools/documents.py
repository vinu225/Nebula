"""
Sovereign AI Workbench — Industrial Document Generators

Generates native Word (.docx), PowerPoint (.pptx), Excel (.xlsx), and PDF files.
Specifically implements Indian PSU / MRPL approval note standards.
"""

from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pptx import Presentation
from pptx.util import Inches as PptxInches
from pptx.util import Pt as PptxPt
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from reportlab.lib.styles import getSampleStyleSheet

from app.core.logging import get_logger

logger = get_logger(__name__)


def generate_approval_note_docx(
    subject: str,
    sections: List[Dict[str, str]],
    file_no: Optional[str] = None,
    department: str = "Inspection & Maintenance Dept",
    author: str = "Executive Engineer",
    competent_authority: str = "Chief General Manager (Technical)",
) -> bytes:
    """Generate a formatted Indian PSU-style Approval Note in DOCX format."""
    doc = Document()

    # Normal style adjustments
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)

    # Header section: Confidential tag and File Reference
    header_p = doc.add_paragraph()
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    conf_run = header_p.add_run("CONFIDENTIAL / INTERNAL\n")
    conf_run.bold = True
    conf_run.font.color.rgb = RGBColor(180, 0, 0)

    now_str = datetime.now(timezone.utc).strftime("%d-%b-%Y")
    ref_run = header_p.add_run(f"Ref: {file_no or 'MRPL/AIW/2026/NOTE-042'}\nDate: {now_str}")
    ref_run.font.size = Pt(9.5)

    # Document Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("MANGALORE REFINERY AND PETROCHEMICALS LIMITED\n")
    title_run.bold = True
    title_run.font.size = Pt(14)
    title_run.font.color.rgb = RGBColor(0, 51, 102)

    sub_title_run = title_p.add_run(f"{department.upper()}\nNOTE FOR APPROVAL\n")
    sub_title_run.bold = True
    sub_title_run.font.size = Pt(12)

    # Subject block
    subj_table = doc.add_table(rows=1, cols=2)
    subj_table.autofit = False
    subj_table.columns[0].width = Inches(1.2)
    subj_table.columns[1].width = Inches(5.3)

    cell_label = subj_table.cell(0, 0)
    cell_label.paragraphs[0].add_run("SUBJECT:").bold = True

    cell_content = subj_table.cell(0, 1)
    cell_content.paragraphs[0].add_run(subject).bold = True

    doc.add_paragraph()  # spacing

    # Body sections
    for sec_idx, sec in enumerate(sections, start=1):
        heading = sec.get("title", f"Section {sec_idx}")
        content = sec.get("content", "")

        h = doc.add_heading(level=2)
        h_run = h.add_run(f"{sec_idx}.0  {heading.upper()}")
        h_run.font.color.rgb = RGBColor(0, 51, 102)
        h_run.font.size = Pt(12)

        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(8)
        p.add_run(content)

    doc.add_paragraph()

    # Sign-off Table
    sign_table = doc.add_table(rows=2, cols=2)
    sign_table.columns[0].width = Inches(3.2)
    sign_table.columns[1].width = Inches(3.2)

    sign_table.cell(0, 0).paragraphs[0].add_run("Initiated by:\n\n\n").bold = True
    sign_table.cell(0, 0).paragraphs[0].add_run(f"({author})\n{department}")

    sign_table.cell(0, 1).paragraphs[0].add_run("Recommended for Approval:\n\n\n").bold = True
    sign_table.cell(0, 1).paragraphs[0].add_run(f"({competent_authority})")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_presentation_pptx(title: str, slides: List[Dict[str, Any]]) -> bytes:
    """Generate a PowerPoint presentation."""
    prs = Presentation()

    # Title slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = title
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = "Sovereign AI Workbench — Industrial Knowledge Work"

    # Content slides
    bullet_layout = prs.slide_layouts[1]
    for s in slides:
        slide = prs.slides.add_slide(bullet_layout)
        slide.shapes.title.text = s.get("title", "Untitled")
        body = slide.shapes.placeholders[1]
        tf = body.text_frame

        points = s.get("points", [])
        if points:
            tf.text = points[0]
            for pt in points[1:]:
                p = tf.add_paragraph()
                p.text = str(pt)
                p.level = 0

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def generate_spreadsheet_xlsx(sheet_title: str, headers: List[str], rows: List[List[Any]]) -> bytes:
    """Generate a styled Excel spreadsheet."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    # Header formatting
    header_fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    for row_data in rows:
        ws.append(row_data)

    # Style rows
    for row in ws.iter_rows(min_row=2, max_row=len(rows) + 1, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_pdf_report(title: str, content_paragraphs: List[str]) -> bytes:
    """Generate a PDF document via ReportLab."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter)
    styles = getSampleStyleSheet()
    story = [
        Paragraph(f"<b>{title}</b>", styles["Title"]),
        Spacer(1, 14),
    ]

    for p in content_paragraphs:
        story.append(Paragraph(p, styles["Normal"]))
        story.append(Spacer(1, 10))

    doc.build(story)
    return buf.getvalue()
